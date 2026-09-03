from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageOps, ImageDraw
import io, shutil, json

ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "sources" / "true-north-decks"
OUT_ROOT = ROOT / "sources" / "true-north-extracted"

DECKS = {
    "investor-ai": SOURCE_DIR / "Investor_AI_source.pptx",
    "boss-ai": SOURCE_DIR / "Boss_AI_source.pptx",
    "topsales-ai": SOURCE_DIR / "TopSales_AI_source.pptx",
}


def picture_candidates(shape):
    found = []
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        found.append(shape)
    elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            found.extend(picture_candidates(child))
    return found


def make_contact_sheet(images, output, title):
    thumb_w, thumb_h = 480, 270
    label_h = 38
    cols = 3
    rows = (len(images) + cols - 1) // cols
    canvas = Image.new("RGB", (cols * thumb_w, 72 + rows * (thumb_h + label_h)), "#eef0eb")
    draw = ImageDraw.Draw(canvas)
    draw.text((24, 22), title, fill="#10251d")
    for i, image_path in enumerate(images):
        img = Image.open(image_path).convert("RGB")
        fitted = ImageOps.contain(img, (thumb_w - 20, thumb_h - 20))
        x = (i % cols) * thumb_w + (thumb_w - fitted.width) // 2
        y = 72 + (i // cols) * (thumb_h + label_h) + 10
        canvas.paste(fitted, (x, y))
        draw.text(((i % cols) * thumb_w + 18, 72 + (i // cols) * (thumb_h + label_h) + thumb_h + 4), f"Slide {i+1:02d}", fill="#315a49")
    canvas.save(output, quality=92)


def main():
    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    manifest = {}
    for slug, pptx_path in DECKS.items():
        prs = Presentation(pptx_path)
        outdir = OUT_ROOT / slug
        outdir.mkdir(parents=True, exist_ok=True)
        slide_files = []
        deck_info = []
        for idx, slide in enumerate(prs.slides, start=1):
            pics = []
            for shape in slide.shapes:
                pics.extend(picture_candidates(shape))
            if not pics:
                raise RuntimeError(f"{slug} slide {idx}: no picture shape found")
            # The decks are image-native. Select the picture with the largest slide area.
            pic = max(pics, key=lambda p: p.width * p.height)
            blob = pic.image.blob
            ext = pic.image.ext.lower().replace("jpeg", "jpg")
            image_path = outdir / f"slide-{idx:02d}.{ext}"
            image_path.write_bytes(blob)
            with Image.open(io.BytesIO(blob)) as im:
                width, height = im.size
            slide_files.append(image_path)
            deck_info.append({
                "slide": idx,
                "file": str(image_path.relative_to(ROOT)),
                "width": width,
                "height": height,
                "picture_shapes": len(pics),
            })
        sheet = OUT_ROOT / f"{slug}-contact-sheet.jpg"
        make_contact_sheet(slide_files, sheet, f"{slug} — {len(slide_files)} slides")
        manifest[slug] = {"slides": deck_info, "contact_sheet": str(sheet.relative_to(ROOT))}
        print(slug, len(slide_files), sheet)
    (OUT_ROOT / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
